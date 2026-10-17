"""Excel export of a run and validation of its outputs.

Typical usage example:

  path, n_cartoons, n_review = export_run(project, "anthropic-claude-opus-5-5")
"""

import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .framework import CODING_DIMENSIONS, STAGE_FILES, STAGES, so_list
from .project import ADJUDICATED, read_json
from .review import reviews
from .validate import Validator


def check_item(project, validator, item_dir):
    outs, res = {}, {}
    for st in STAGES:
        f = item_dir / STAGE_FILES[st]
        invalid = (item_dir / STAGE_FILES[st].replace(".json", ".invalid.json")).is_file()
        if f.is_file():
            obj = read_json(f)
            outs[st] = obj
            if st == 4 and not validator.fw.has_coding:
                res[st] = (obj, [], ["no codebook.yaml to check Stage 4 against"], invalid)
                continue
            errs, warns = validator.stage(st, obj, item_dir.name, outs)
            res[st] = (obj, errs, warns, invalid)
        elif invalid:
            res[st] = (None, [], [], True)
    return res


def validate_run(project, run):
    v = Validator(project.framework())
    for d in project.item_dirs(run):
        for st, (obj, errs, warns, _) in check_item(project, v, d).items():
            if obj is not None:
                yield d.name, st, errs, warns


def _sheet(wb, title, rows, cols):
    ws = wb.create_sheet(title)
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([_cell(r.get(k)) for k in cols])
    ws.freeze_panes = "B2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for i, k in enumerate(cols, 1):
        width = max([len(str(k))] + [min(len(str(r.get(k) or "")), 60) for r in rows[:200]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 62)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=isinstance(c.value, str) and len(c.value) > 60)
    return ws


def _cell(v):
    if isinstance(v, (list, dict)):
        return json.dumps(v, ensure_ascii=False)
    if isinstance(v, bool):
        return "yes" if v else "no"
    return v


def export_run(project, run, path=None):
    fw = project.framework()
    v = Validator(fw)
    items = {i["cartoon_id"]: i for i in project.items(check_dates=False)}
    meta_cols = project.columns(list(items.values()))
    cartoons, sos, orients, targets, reals, presups, review = [], [], [], [], [], [], []
    for d in project.item_dirs(run):
        cid = d.name
        res = check_item(project, v, d)
        o = {st: r[0] for st, r in res.items()}
        for st in STAGES:
            obj, errs, warns, invalid = res.get(st, (None, [], [], False))
            status = obj.get("annotation_status") if obj else None
            if invalid or errs or warns or status == "needs_review" or (obj is None and st in (1, 2, 3)):
                review.append({"cartoon_id": cid, "stage": st,
                               "state": "invalid" if invalid and obj is None else ("missing" if obj is None else status),
                               "review_reasons": " | ".join(obj.get("review_reasons", [])) if obj else "",
                               "validation_errors": " | ".join(errs), "validation_warnings": " | ".join(warns),
                               "difficulty_notes": (obj or {}).get("difficulty_notes")})
        row = {**{k: items.get(cid, {}).get(k) for k in meta_cols}, "cartoon_id": cid,
               "genre": project.settings.get("material") or "editorial cartoons"}
        if run != ADJUDICATED:
            for st, rec in reviews(project, cid).items():
                if rec.get("run") == run:
                    row[f"stage{st}_review"] = rec["status"]
        s1, s2, s3, s4 = o.get(1), o.get(2), o.get(3), o.get(4)
        if s1:
            sr = s1["shared_reading"]
            row.update(stage1_status=s1["annotation_status"], scene_description=sr["scene_description"],
                       text_status=sr["text_status"], text_transcription=sr["text_transcription"],
                       n_references=len(sr["references"]), n_so=len(so_list(s1)))
            coded = {x["so_id"]: x for x in (s4 or {}).get("oppositions", [])}
            for s in so_list(s1):
                c = coded.get(s["id"], {})
                sos.append({"cartoon_id": cid, "so_id": s["id"], "is_primary": s["id"] == "so_1",
                            **{k: s[k] for k in ("concrete_opposition", "intermediate_label", "abstract_code",
                                                 "abstract_alternative")},
                            "opposition_category": c.get("category"),
                            "opposition_proposed": c.get("proposed_category"),
                            "confidence": s["confidence"], "justification": s["justification"]})
        if s2:
            ns = s2["narrative_strategy"]
            row.update(stage2_status=s2["annotation_status"], si_summary=s2["situation"]["summary"],
                       si_modality=s2["situation"]["modality"], layout=ns["layout"], panel_count=ns["panel_count"],
                       devices=";".join(x["code"] for x in ns["devices"]),
                       humour_distribution=ns["humour_distribution"], n_targets=len(s2["targets"]))
            for x in s2["orientations"]:
                orients.append({"cartoon_id": cid, **x})
            coded = {t["entity"]: t for t in (s4 or {}).get("targets", [])}
            for t in s2["targets"]:
                c = coded.get(t["entity"], {})
                targets.append({"cartoon_id": cid, **t, "category": c.get("category"),
                                "proposed_category": c.get("proposed_category"),
                                "category_confidence": c.get("confidence")})
        if s3:
            th = s3["theme"]
            row.update(stage3_status=s3["annotation_status"], theme_primary=th["primary"],
                       theme_secondary=th["secondary"], theme_confidence=th["confidence"])
            for r in s3["realizations"]:
                reals.append({"cartoon_id": cid, "so_id": r["so_id"], "locus": r["locus"],
                              "interaction_dynamic": r["interaction_dynamic"],
                              "interaction_alternative": r["interaction_alternative"],
                              "semiotic_mechanisms": ";".join(m["code"] for m in r["semiotic_mechanisms"]),
                              "n_metaphors": len(r["metaphors"]), "metaphors": r["metaphors"] or None,
                              "confidence": r["confidence"], "note": r["note"]})
            for p in s3["presuppositions"]:
                presups.append({"cartoon_id": cid, **p})
        if s4:
            row["stage4_status"] = s4["annotation_status"]
            th4 = s4.get("theme")
            if th4:
                for slot in ("primary", "secondary"):
                    row[f"theme_{slot}_category"] = th4[slot]["category"] if th4[slot] else None
                    row[f"theme_{slot}_proposed"] = th4[slot]["proposed_category"] if th4[slot] else None
        cartoons.append(row)

    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "cartoons", cartoons, meta_cols + ["genre", 
        "stage1_status", "stage2_status", "stage3_status", "stage4_status", "stage1_review", "stage2_review",
        "stage3_review", "stage4_review", "scene_description", "text_status",
        "text_transcription", "n_references", "n_so", "si_summary", "si_modality", "layout", "panel_count",
        "devices", "humour_distribution", "n_targets", "theme_primary", "theme_secondary", "theme_confidence",
        "theme_primary_category", "theme_primary_proposed", "theme_secondary_category", "theme_secondary_proposed"])
    _sheet(wb, "oppositions", sos, ["cartoon_id", "so_id", "is_primary", "concrete_opposition", "intermediate_label",
                                    "opposition_category", "opposition_proposed", "abstract_code",
                                    "abstract_alternative", "confidence", "justification"])
    _sheet(wb, "orientations", orients, ["cartoon_id", "id", "orientation", "referent", "aggression", "basis",
                                         "confidence"])
    _sheet(wb, "targets", targets, ["cartoon_id", "orientation_id", "role", "entity", "target_type",
                                    "criticized_attribute", "basis", "confidence", "category", "proposed_category",
                                    "category_confidence"])
    _sheet(wb, "realizations", reals, ["cartoon_id", "so_id", "locus", "interaction_dynamic",
                                       "interaction_alternative", "semiotic_mechanisms", "n_metaphors", "metaphors",
                                       "confidence", "note"])
    _sheet(wb, "presuppositions", presups, ["cartoon_id", "category", "specific_knowledge", "necessity", "basis",
                                            "confidence"])
    _sheet(wb, "review_queue", review, ["cartoon_id", "stage", "state", "review_reasons", "validation_errors",
                                        "validation_warnings", "difficulty_notes"])
    if fw.codebook:
        cb_rows = [{"dimension": dim, "code": c["code"], "definition": c["definition"]}
                   for dim, name in CODING_DIMENSIONS.items() if name in fw.cb for c in fw.cb[name]]
        _sheet(wb, f"codebook {fw.codebook['version']}"[:31], cb_rows, ["dimension", "code", "definition"])
    project.exports_dir.mkdir(parents=True, exist_ok=True)
    path = path or project.exports_dir / f"{run}.xlsx"
    wb.save(path)
    return path, len(cartoons), len(review)
